"""Генерация и разбор Excel-отчёта по результатам пайплайна.

build_excel_report: PipelineResult -> xlsx (4 листа).
Листы и заголовки — константы: parse_edited_report (ниже в этом же файле,
Task 7) читает их обратно для пересоздания заказа.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import TYPE_CHECKING, Any, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from process_specification_table import detect_product_type

if TYPE_CHECKING:
    from bitrix_bot.pipeline import PipelineResult

SHEET_LOADED = "Загружено"
SHEET_SKIPPED = "Пропущено"
SHEET_TRADING = "Перекупное"
SHEET_ERRORS = "Ошибки 1С"

TRADING_REASON_PREFIXES = ("Покупная позиция", "Покупная арматура")
ORDER_NUMBER_LABEL = "Номер заказа"

# Product types treated as resold equipment (покупное оборудование).
# Источник истины — здесь (web_app.py не импортируется: это Streamlit-приложение).
EQUIPMENT_PTYPES = {
    "diffuser",
    "ksd",
    "grille",
    "silencer",
    "damper",
    "shutter",
    "filter",
    "throttle",
    "roof_cap",
    "fan",
}

LOADED_HEADERS = ["Артикул", "A", "B", "D", "Кол-во", "Материал", "Толщина",
                  "Соед. 0", "Соед. 1", "Соед. 2", "Соед. 3",
                  "Система", "Комментарий"]
SKIPPED_HEADERS = ["Наименование", "Размер", "Кол-во", "Ед.", "Материал",
                   "Толщина", "Причина", "Включить в заказ"]
TRADING_HEADERS = ["Наименование", "Категория", "Размер", "Материал", "Толщина",
                   "Кол-во", "Ед.", "Производитель", "Модель", "Включить в заказ"]


def is_trading_skip(skip: dict) -> bool:
    """Перекупное (покупное) оборудование: скипы по причине «Покупная …»,
    строки автомаппинга оборудования (raw_name) и позиции, чей тип
    распознаётся как покупное оборудование (EQUIPMENT_PTYPES)."""
    if skip.get("raw_name"):
        return True
    reason = str(skip.get("reason", ""))
    if reason.startswith(TRADING_REASON_PREFIXES):
        return True
    name = skip.get("name") or skip.get("raw_name") or ""
    if name:
        ptype = detect_product_type(name)
        if ptype in EQUIPMENT_PTYPES:
            return True
    return False


def _write_sheet(ws, headers: List[str], rows: List[List[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        width = len(str(h))
        for row in rows:
            width = max(width, len(str(row[i - 1])) if row[i - 1] is not None else 0)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)


def _loaded_row(p: dict) -> List[Any]:
    dims = p.get("params") or {}
    return [
        p.get("article", ""), dims.get("A0"), dims.get("B0"), dims.get("D0"),
        p.get("quantity"), p.get("material_code", ""), p.get("thickness"),
        p.get("connection_0", ""), p.get("connection_1", ""),
        p.get("connection_2", ""), p.get("connection_3", ""),
        p.get("system", ""), p.get("comment", ""),
    ]


def _skipped_row(s: dict) -> List[Any]:
    return [s.get("name", ""), s.get("size", ""), s.get("quantity"),
            s.get("unit", ""), s.get("material", ""), s.get("thickness"),
            s.get("reason", ""), ""]


def _trading_row(s: dict) -> List[Any]:
    name = s.get("name") or s.get("raw_name") or ""
    return [name, detect_product_type(name) or "",
            s.get("size") or s.get("model") or "",
            s.get("material", ""), s.get("thickness"),
            s.get("quantity"), s.get("unit", "шт"),
            s.get("manufacturer", ""), s.get("model", ""), ""]


def build_excel_report(res: "PipelineResult") -> bytes:
    """Собрать xlsx-отчёт: Загружено / Пропущено / Перекупное / Ошибки 1С."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_LOADED
    _write_sheet(ws, LOADED_HEADERS, [_loaded_row(p) for p in res.loaded])

    trading = [s for s in res.skipped if is_trading_skip(s)]
    skipped = [s for s in res.skipped if not is_trading_skip(s)]

    _write_sheet(wb.create_sheet(SHEET_SKIPPED), SKIPPED_HEADERS,
                 [_skipped_row(s) for s in skipped])
    _write_sheet(wb.create_sheet(SHEET_TRADING), TRADING_HEADERS,
                 [_trading_row(s) for s in trading])

    ws = wb.create_sheet(SHEET_ERRORS)
    summary = [
        (ORDER_NUMBER_LABEL, res.order_number or ""),
        ("Файл", res.file_name),
        ("Загружено", len(res.loaded)),
        ("Пропущено", len(skipped)),
        ("Перекупное", len(trading)),
        ("Ошибок 1С", len(res.errors_1c)),
        ("Предупреждений 1С", len(res.warnings_1c)),
    ]
    ws.append(["Показатель", "Значение"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for label, value in summary:
        ws.append([label, value])
    ws.append([])
    ws.append(["Тип", "Текст"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for e in res.errors_1c:
        ws.append(["Ошибка", e])
    for w in res.warnings_1c:
        ws.append(["Предупреждение", w])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class EditedReportError(ValueError):
    """Отредактированный отчёт не читается (формат/значения)."""


INCLUDE_YES = {"да", "yes", "1", "+"}


@dataclass
class EditedReport:
    loaded_rows: List[dict]
    include_rows: List[dict]
    skipped_rows: List[dict]
    replaced_order: Optional[str] = None


def _s(value) -> str:
    return "" if value is None else str(value).strip()


def _num(value, ctx: str) -> float:
    try:
        return float(str(value).replace(",", ".").replace(" ", ""))
    except (TypeError, ValueError):
        raise EditedReportError(f"{ctx}: ожидалось число, получено {value!r}")


def _thickness_or_default(value) -> float:
    """Толщина из колонки «Толщина»; пустая/битая → дефолт 0.8."""
    try:
        t = float(str(value).replace(",", ".").replace(" ", ""))
        return t if t > 0 else 0.8
    except (TypeError, ValueError):
        return 0.8


def parse_edited_report(content: bytes) -> EditedReport:
    """Прочитать отредактированный отчёт обратно в позиции для пересоздания заказа."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise EditedReportError(f"Не удалось прочитать xlsx: {e}")

    for name in (SHEET_LOADED, SHEET_ERRORS):
        if name not in wb.sheetnames:
            raise EditedReportError(f"В файле нет листа «{name}» — это не отчёт системы")

    replaced_order = None
    loaded_rows: List[dict] = []
    include_rows: List[dict] = []
    skipped_rows: List[dict] = []

    # --- Загружено ---
    ws = wb[SHEET_LOADED]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for idx, r in enumerate(rows, start=2):
        if not any(c is not None and str(c).strip() for c in r):
            continue  # удалённая пользователем строка — пропускаем молча
        ctx = f"Лист «{SHEET_LOADED}», строка {idx}"
        article = _s(r[0])
        if not article:
            raise EditedReportError(f"{ctx}: пустой артикул")
        params = {}
        for key, cell in (("A0", r[1]), ("B0", r[2]), ("D0", r[3])):
            if cell is not None and str(cell).strip():
                params[key] = _num(cell, ctx)
        if not params:
            raise EditedReportError(f"{ctx}: нет ни одного размера (A/B/D)")
        quantity = _num(r[4], ctx)
        if quantity <= 0:
            raise EditedReportError(f"{ctx}: количество должно быть > 0")
        q = int(quantity) if float(quantity).is_integer() else quantity
        thickness = _num(r[6], ctx)
        if thickness <= 0:
            raise EditedReportError(f"{ctx}: толщина должна быть > 0")
        loaded_rows.append({
            "article": article,
            "params": params,
            "quantity": q,
            "material_code": _s(r[5]) or "1",
            "thickness": thickness,
            "connection_0": _s(r[7]), "connection_1": _s(r[8]),
            "connection_2": _s(r[9]), "connection_3": _s(r[10]),
            "system": _s(r[11]),
            "comment": _s(r[12]),
        })

    # --- Пропущено / Перекупное: индексы колонок из заголовков листа ---
    for sheet_name, headers in ((SHEET_SKIPPED, SKIPPED_HEADERS),
                                (SHEET_TRADING, TRADING_HEADERS)):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        actual = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        def col(title: str) -> int:
            try:
                return actual.index(title)
            except ValueError:
                raise EditedReportError(
                    f"Лист «{sheet_name}»: нет колонки «{title}» — файл изменён вне отчёта")
        i_size, i_qty = col("Размер"), col("Кол-во")
        i_unit, i_mat = col("Ед."), col("Материал")
        i_inc = col("Включить в заказ")
        i_thick = col("Толщина")
        for idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(c is not None and str(c).strip() for c in r):
                continue
            name = _s(r[0])
            if not name:
                continue
            item = {
                "name": name,
                "size": _s(r[i_size]),
                "unit": _s(r[i_unit]) or "шт",
                "quantity": 1.0,
                "material": _s(r[i_mat]) or "оцинкованная",
                "thickness": _thickness_or_default(r[i_thick]),
            }
            q_raw = r[i_qty]
            if q_raw is not None and str(q_raw).strip():
                try:
                    item["quantity"] = float(str(q_raw).replace(",", "."))
                except ValueError:
                    pass
            skipped_rows.append(item)
            if _s(r[i_inc]).lower() in INCLUDE_YES:
                # служебный runtime-маркер: включённая позиция не должна
                # оставаться в skipped при пересоздании заказа
                item["_include"] = True
                include_rows.append(item)

    # --- Ошибки 1С: номер заменяемого заказа из сводки (колонка A) ---
    ws = wb[SHEET_ERRORS]
    for r in ws.iter_rows(values_only=True):
        if r and _s(r[0]) == ORDER_NUMBER_LABEL and len(r) > 1:
            replaced_order = _s(r[1]) or None
            break

    # Отказ только если пусто И включённых нет: пересоздание возможно
    # и из одних включённых позиций (аналоги подбираются заново).
    if not loaded_rows and not include_rows:
        raise EditedReportError(
            "Лист «Загружено» пуст и позиции на включение отсутствуют — "
            "нечего пересоздавать"
        )

    return EditedReport(
        loaded_rows=loaded_rows,
        include_rows=include_rows,
        skipped_rows=skipped_rows,
        replaced_order=replaced_order,
    )
