"""Тесты пайплайна PDF → 1С (MCP execute_code замокан)."""
import io

import httpx
import pytest

import bitrix_bot.pipeline as pl
from bitrix_bot.pipeline import PipelineResult
from tools.as_order_loader.build_execute_payload import build_payload

SAMPLE_1C_OK = (
    "ЗАКАЗ 000000860 | строк=2 | ошибок=0 | предупр=1"
    " | Строка 1 (1-2-1): цена 0 — проверьте прайс"
    " ## 1|1-2-1|t=0.8|мат=Рулон оц.(08ПС)0.80|цена=0|S=1.2|n=6"
    " ## 2|4-2-3|t=0.8|мат=Рулон оц.(08ПС)0.80|цена=150|S=0.4|n=2"
)

SAMPLE_1C_ERRORS = (
    "ЗАКАЗ 000000861 | строк=2 | ошибок=1"
    " | Строка 1: не найден продукт \"9-9-9\""
    " | предупр=0"
)


class _Resp:
    def __init__(self, data):
        self.status_code = 200
        self._data = data

    def json(self):
        return self._data

    def raise_for_status(self):
        return None


def test_build_payload_order_comment_substituted():
    payload = build_payload([{"article": "1-2-1"}], order_comment="Задача №42: Шипиловский")
    assert "Задача №42: Шипиловский" in payload["code"]
    # по умолчанию — прежний комментарий
    payload_def = build_payload([{"article": "1-2-1"}])
    assert "Загрузка из JSON (execute_code, build_execute_payload)" in payload_def["code"]


def test_build_payload_escapes_order_comment_quotes():
    payload = build_payload(
        [{"article": "1-2-1"}], order_comment='Задача №1: сказал "ура"'
    )
    # кавычки 1С в строковом литерале удваиваются
    assert 'сказал ""ура""' in payload["code"]


def test_parse_1c_result_errors_joined_single_segment():
    # BSL склеивает ошибки/предупреждения через СтрСоединить(..., "; ") —
    # всё сообщение приходит ОДНИМ сегментом после "ошибок=N"
    text = (
        'ЗАКАЗ 000001 | строк=2 | ошибок=2'
        ' | не найден продукт "9-9-9"; не найден материал "X"'
        ' | предупр=1 | Строка 1: цена 0 — проверьте прайс'
    )
    out = pl.parse_1c_result(text)
    assert out["order_number"] == "000001"
    assert out["errors"] == ['не найден продукт "9-9-9"', 'не найден материал "X"']
    assert out["warnings"] == ["Строка 1: цена 0 — проверьте прайс"]


def test_load_order_parses_ok(monkeypatch):
    monkeypatch.setattr(
        httpx, "post",
        lambda url, json=None, timeout=None: _Resp({"result": SAMPLE_1C_OK}),
    )
    out = pl.load_order_to_1c([{"article": "1-2-1"}], "http://x/api/execute_code", "c")
    assert out["order_number"] == "000000860"
    assert out["errors"] == []
    assert out["warnings"] == ["Строка 1 (1-2-1): цена 0 — проверьте прайс"]


def test_load_order_parses_errors(monkeypatch):
    monkeypatch.setattr(
        httpx, "post",
        lambda url, json=None, timeout=None: _Resp({"result": SAMPLE_1C_ERRORS}),
    )
    out = pl.load_order_to_1c([{"article": "9-9-9"}], "http://x/api/execute_code", "c")
    assert out["order_number"] == "000000861"
    assert out["errors"] == ['Строка 1: не найден продукт "9-9-9"']
    assert out["warnings"] == []


def test_load_order_parses_native_envelope(monkeypatch):
    # нативный HTTP-API MCPToolkit (127.0.0.1:6005): {"success": true, "data": "..."}
    # — раньше str(dict) попадал в парсер: номер заказа терялся, а в ошибки
    # улетал сегмент "предупр=N" (регрессия после перехода с tunnel-прокси).
    monkeypatch.setattr(
        httpx, "post",
        lambda url, json=None, timeout=None: _Resp(
            {"success": True, "data": SAMPLE_1C_OK}
        ),
    )
    out = pl.load_order_to_1c([{"article": "1-2-1"}], "http://x/api/execute_code", "c")
    assert out["order_number"] == "000000860"
    assert out["errors"] == []
    assert out["warnings"] == ["Строка 1 (1-2-1): цена 0 — проверьте прайс"]


def test_load_order_native_envelope_1c_error(monkeypatch):
    # ошибка 1С внутри успешного HTTP-ответа: {"success": true, "data": "ЗАКАЗ ... ошибок=1 ..."}
    monkeypatch.setattr(
        httpx, "post",
        lambda url, json=None, timeout=None: _Resp(
            {"success": True, "data": SAMPLE_1C_ERRORS}
        ),
    )
    out = pl.load_order_to_1c([{"article": "9-9-9"}], "http://x/api/execute_code", "c")
    assert out["order_number"] == "000000861"
    assert out["errors"] == ['Строка 1: не найден продукт "9-9-9"']


def test_load_order_native_envelope_failure(monkeypatch):
    # {"success": false, "error": "..."} — исключение с текстом 1С, не молчаливый dict-repr
    monkeypatch.setattr(
        httpx, "post",
        lambda url, json=None, timeout=None: _Resp(
            {"success": False, "error": "Ошибка компиляции"}
        ),
    )
    with pytest.raises(RuntimeError, match="Ошибка компиляции"):
        pl.load_order_to_1c([{"article": "1-2-1"}], "http://x/api/execute_code", "c")


def test_run_pipeline_no_positions(monkeypatch):
    monkeypatch.setattr(
        pl, "process_pdf_to_positions",
        lambda b: ([], [{"name": "x", "reason": "bad"}]),
    )
    res = pl.run_pipeline(b"pdf", "f.pdf", "c", "http://x")
    assert res.order_number is None
    assert res.loaded == []
    assert len(res.skipped) == 1
    assert "не распознал" in res.raw_text


def test_run_pipeline_end_to_end_mocked(monkeypatch):
    monkeypatch.setattr(
        pl, "process_pdf_to_positions",
        lambda b: (
            [{"article": "1-2-1", "quantity": 6, "comment": "", "thickness": 0.8, "params": {}}],
            [{"name": "Клапан", "size": "200", "quantity": "1", "reason": "нет маппинга"}],
        ),
    )
    monkeypatch.setattr(
        httpx, "post",
        lambda url, json=None, timeout=None: _Resp({"result": SAMPLE_1C_OK}),
    )
    res = pl.run_pipeline(b"pdf", "spec.pdf", "Задача №42", "http://x")
    assert res.order_number == "000000860"
    assert len(res.loaded) == 1
    assert res.skipped[0]["name"] == "Клапан"
    assert res.warnings_1c == ["Строка 1 (1-2-1): цена 0 — проверьте прайс"]


def test_recreate_order_from_report(monkeypatch):
    import report_xlsx

    res_in = PipelineResult(
        file_name="spec.pdf", order_number="839",
        loaded=[{"article": "1-1-1", "params": {"A0": 300, "B0": 200},
                 "quantity": 2, "material_code": "1", "thickness": 0.8,
                 "connection_0": "6", "connection_1": "2",
                 "connection_2": "0", "connection_3": "0",
                 "system": "", "comment": "Воздуховод 300x200"}],
        skipped=[{"name": "Гибкий воздуховод Ф125", "size": "Ф125", "unit": "м",
                  "quantity": 10, "material": "оцинкованная", "thickness": 0.8,
                  "reason": "Покупная позиция (…) — завод не производит"}],
    )
    xlsx_bytes = report_xlsx.build_excel_report(res_in)

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    wb[report_xlsx.SHEET_LOADED].cell(row=2, column=7, value=1.0)  # толщина 1.0
    buf = io.BytesIO()
    wb.save(buf)

    captured = {}

    def fake_load(positions, execute_url, order_comment, timeout=280.0):
        captured["positions"] = positions
        captured["comment"] = order_comment
        return {"order_number": "840", "errors": [], "warnings": [],
                "raw": "ЗАКАЗ 840 | строк=1 | ошибок=0 | предупр=0"}

    monkeypatch.setattr(pl, "load_order_to_1c", fake_load)

    out = pl.recreate_order_from_report(
        buf.getvalue(), "http://x", base_comment="Задача Битрикс24 №42: Т")

    assert captured["comment"] == "Задача Битрикс24 №42: Т | Заменяет заказ №839 (исправлено из отчёта)"
    assert captured["positions"][0]["thickness"] == 1.0
    assert captured["positions"][0]["material"] == "Рулон оц.(08ПС)1.00"
    assert out.order_number == "840"
    assert len(out.loaded) == 1
    assert len(out.skipped) == 1  # перекупное осталось пропущенным


def _report_with_include(skipped_row, loaded_rows=None):
    """Отчёт с одной позицией «Пропущено», у которой стоит «Включить в заказ» = да."""
    import report_xlsx
    from openpyxl import load_workbook

    res_in = PipelineResult(
        file_name="spec.pdf", order_number="839",
        loaded=loaded_rows or [], skipped=[skipped_row],
    )
    content = report_xlsx.build_excel_report(res_in)
    wb = load_workbook(io.BytesIO(content))
    for sheet in (report_xlsx.SHEET_SKIPPED, report_xlsx.SHEET_TRADING):
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        inc_col = headers.index("Включить в заказ") + 1
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == skipped_row["name"]:
                ws.cell(row=row, column=inc_col, value="да")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fake_1c_load(positions, execute_url, order_comment, timeout=280.0):
    return {"order_number": "840", "errors": [], "warnings": [],
            "raw": "ЗАКАЗ 840 | строк=1 | ошибок=0 | предупр=0"}


def test_recreate_include_success_loaded_once_not_skipped(monkeypatch):
    # I-2: включённая позиция с подобранным аналогом — ровно один раз в loaded,
    # ни одного её упоминания в skipped.
    content = _report_with_include(
        {"name": "Воздуховод прямошовный 300x200", "size": "300x200", "unit": "м",
         "quantity": 5, "material": "оцинкованная", "thickness": 0.8,
         "reason": "Не удалось распознать артикул"},
    )
    monkeypatch.setattr(pl, "load_order_to_1c", _fake_1c_load)

    out = pl.recreate_order_from_report(content, "http://x")

    hits = [p for p in out.loaded if p.get("article") == "1-2-1"]
    assert len(hits) == 1
    assert not any("Воздуховод прямошовный" in str(s.get("name", ""))
                   for s in out.skipped)


def test_recreate_include_unparsable_single_skip_with_reason(monkeypatch):
    # I-2: включённая позиция без аналога — ровно один раз в skipped, с причиной
    # (а не две строки: без reason из skipped_rows + с reason из extra_skipped).
    content = _report_with_include(
        {"name": "Абракадабра без размера", "size": "", "unit": "шт",
         "quantity": 2, "material": "оцинкованная", "thickness": 0.8,
         "reason": "Не удалось распознать артикул"},
        loaded_rows=[{"article": "1-1-1", "params": {"A0": 300, "B0": 200},
                     "quantity": 2, "material_code": "1", "thickness": 0.8}],
    )
    monkeypatch.setattr(pl, "load_order_to_1c", _fake_1c_load)

    out = pl.recreate_order_from_report(content, "http://x")

    matches = [s for s in out.skipped
               if "Абракадабра" in str(s.get("name", ""))]
    assert len(matches) == 1
    assert matches[0].get("reason")
