import io

import pytest
from openpyxl import load_workbook

from bitrix_bot.pipeline import PipelineResult
from report_xlsx import (
    SHEET_ERRORS, SHEET_LOADED, SHEET_SKIPPED, SHEET_TRADING,
    EditedReportError, build_excel_report, is_trading_skip, parse_edited_report,
)


def _sample_result() -> PipelineResult:
    return PipelineResult(
        file_name="spec.pdf",
        order_number="839",
        loaded=[
            {"article": "1-1-1", "params": {"A0": 300, "B0": 200}, "quantity": 2,
             "material_code": "1", "thickness": 0.8,
             "connection_0": "6", "connection_1": "2",
             "connection_2": "0", "connection_3": "0",
             "system": "П1", "comment": "Воздуховод 300x200"},
        ],
        skipped=[
            {"name": "Гибкий воздуховод Ф125", "size": "Ф125", "unit": "м",
             "quantity": 10, "material": "оцинкованная", "thickness": 0.8,
             "reason": "Покупная позиция (гибкий воздуховод/трубопровод/изоляция) — завод не производит"},
            {"name": "Клапан ОЗ-60-НО-400*200", "size": "400x200", "unit": "шт",
             "quantity": 1, "material": "оцинкованная", "thickness": 0.8,
             "reason": "Покупная арматура (брендовый клапан/шумоглушитель) — завод не производит"},
            {"name": "Неизвестная деталь", "size": "100x100", "unit": "шт",
             "quantity": 3, "material": "оцинкованная", "thickness": 0.8,
             "reason": "Не удалось распознать артикул"},
            {"name": "Диффузор без артикула", "size": "300", "unit": "шт",
             "quantity": 2, "reason": "Диффузор без артикула в каталоге"},
            {"raw_name": "Диффузор SR-P 300", "model": "SR-P-300",
             "reason": "Оборудование не производится"},
        ],
        errors_1c=["Нет цены у артикула 1-1-1"],
        warnings_1c=["Проверьте шину"],
    )


def test_sheets_and_split():
    wb = load_workbook(io.BytesIO(build_excel_report(_sample_result())))
    assert wb.sheetnames == [SHEET_LOADED, SHEET_SKIPPED, SHEET_TRADING, SHEET_ERRORS]

    loaded_ws = wb[SHEET_LOADED]
    headers = [c.value for c in loaded_ws[1]]
    assert headers[:6] == ["Артикул", "A", "B", "D", "Кол-во", "Материал"]
    assert "Соед. 0" in headers and "Соед. 3" in headers
    row = [c.value for c in loaded_ws[2]]
    assert row[0] == "1-1-1" and row[1] == 300 and row[2] == 200

    skipped_ws = wb[SHEET_SKIPPED]
    names = [r[0] for r in skipped_ws.iter_rows(min_row=2, values_only=True)]
    assert names == ["Неизвестная деталь"]  # только не-перекупное

    trading_ws = wb[SHEET_TRADING]
    tnames = [r[0] for r in trading_ws.iter_rows(min_row=2, values_only=True)]
    assert "Гибкий воздуховод Ф125" in tnames
    assert "Диффузор SR-P 300" in tnames  # raw_name-строка тоже перекупное
    assert "Диффузор без артикула" in tnames  # по категории оборудования, без «Покупная …»
    theaders = [c.value for c in trading_ws[1]]
    assert "Материал" in theaders and "Толщина" in theaders and "Включить в заказ" in theaders

    errors_ws = wb[SHEET_ERRORS]
    labels = [r[0] for r in errors_ws.iter_rows(values_only=True) if r[0]]
    assert "Номер заказа" in labels
    texts = [r[-1] for r in errors_ws.iter_rows(values_only=True) if r and r[0] in ("Ошибка", "Предупреждение")]
    assert "Нет цены у артикула 1-1-1" in texts


def test_is_trading_skip():
    assert is_trading_skip({"reason": "Покупная позиция (…) — завод не производит"})
    assert is_trading_skip({"reason": "Покупная арматура (…) — завод не производит"})
    assert is_trading_skip({"raw_name": "Диффузор", "reason": "Оборудование не производится"})
    assert not is_trading_skip({"reason": "Не удалось распознать артикул"})


def test_is_trading_skip_by_equipment_ptype():
    # Диффузор — категория оборудования, даже если reason не «Покупная …»
    assert is_trading_skip({"name": "Диффузор SR-P 300",
                            "reason": "Диффузор без артикула в каталоге"})
    # «Неизвестная деталь» не распознаётся как оборудование → не перекупное
    assert not is_trading_skip({"name": "Неизвестная деталь",
                                "reason": "Не удалось распознать артикул"})


def test_parse_edited_report_roundtrip():
    res = _sample_result()
    content = build_excel_report(res)

    # «Правка»: толщина 1.0, добавим «да» у перекупной позиции
    wb = load_workbook(io.BytesIO(content))
    wb[SHEET_LOADED].cell(row=2, column=7, value=1.0)  # Толщина
    wb[SHEET_TRADING].cell(row=2, column=10, value="да")  # Включить в заказ
    buf = io.BytesIO()
    wb.save(buf)

    edited = parse_edited_report(buf.getvalue())
    assert edited.replaced_order == "839"
    assert len(edited.loaded_rows) == 1
    row = edited.loaded_rows[0]
    assert row["article"] == "1-1-1"
    assert row["params"] == {"A0": 300, "B0": 200}
    assert row["thickness"] == 1.0
    assert row["material_code"] == "1"
    assert row["connection_0"] == "6"
    assert len(edited.include_rows) == 1
    assert edited.include_rows[0]["name"] == "Гибкий воздуховод Ф125"
    assert len(edited.skipped_rows) == 5  # все пропущенные/перекупные


def test_parse_edited_report_include_from_skipped_sheet():
    # «Включить в заказ» = «да» у строки «Неизвестная деталь» (лист «Пропущено»)
    content = build_excel_report(_sample_result())
    wb = load_workbook(io.BytesIO(content))
    wb[SHEET_SKIPPED].cell(row=2, column=8, value="да")  # Включить в заказ
    buf = io.BytesIO()
    wb.save(buf)

    edited = parse_edited_report(buf.getvalue())
    assert any(r["size"] == "100x100" for r in edited.include_rows)
    assert edited.include_rows[0]["quantity"] == 3


def test_parse_edited_report_rejects_garbage():
    with pytest.raises(EditedReportError):
        parse_edited_report(b"not an xlsx")


def test_parse_edited_report_empty_loaded():
    # loaded пуст → лист «Загружено» содержит только заголовок
    res = PipelineResult(
        file_name="spec.pdf",
        order_number="839",
        loaded=[],
        skipped=[{"name": "Неизвестная деталь", "size": "100x100", "unit": "шт",
                  "quantity": 3, "material": "оцинкованная", "thickness": 0.8,
                  "reason": "Не удалось распознать артикул"}],
        errors_1c=[],
        warnings_1c=[],
    )
    with pytest.raises(EditedReportError, match="нечего пересоздавать"):
        parse_edited_report(build_excel_report(res))
