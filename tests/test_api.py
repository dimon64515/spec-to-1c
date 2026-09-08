import pandas as pd

from api import (
    ProcessResult,
    read_csv_or_excel_bytes,
    process_specification_file,
)


def test_process_specification_file_csv_round_duct():
    csv = (
        "Наименование;Размер;Ед;Количество;Материал;Толщина\n"
        "Воздуховод;160;м;10;оцинкованная;0.8\n"
    )
    result = process_specification_file(
        csv.encode("utf-8-sig"),
        "spec.csv",
        header={"order_name": "TEST"},
    )
    assert isinstance(result, ProcessResult)
    assert "1-1-2" in result.xml
    assert "D00160" in result.xml
    assert result.skipped == []


def test_read_csv_or_excel_bytes_parses_csv():
    csv = "name;size;unit;quantity\nВоздуховод;160;м;10\n"
    df = read_csv_or_excel_bytes(csv.encode("utf-8-sig"), "spec.csv")
    assert isinstance(df, pd.DataFrame)
    assert len(df) == 1
    assert df.iloc[0]["name"] == "Воздуховод"


def _make_project_spec_xlsx(tmp_path):
    """Минимальный файл проектной спецификации (разметка FineReader)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append([
        "Позиция",
        "Наименование и техническая характеристика",
        "Тип, марка, обозначение документа, опросного листа",
        "Код оборудования, изделия, материала",
        "Завод-изготовитель",
        "Единица измерения",
        "Количество",
    ])
    ws.append(["1", "2", "3", "4", "5", "6", "7"])
    ws.append(["1", "Отвод прямоугольного воздуховода 300x200", "Россия", "", "", "шт", "3"])
    p = tmp_path / "spec.xlsx"
    wb.save(p)
    return p


def test_process_specification_file_project_spec_xlsx_uses_project_parser(tmp_path):
    """Проектная спецификация (FineReader Excel) не должна падать в generic CSV-путь."""
    p = _make_project_spec_xlsx(tmp_path)
    result = process_specification_file(p.read_bytes(), "spec.xlsx")
    assert "<productRow>" in result.xml
    assert "<article>2-2-2</article>" in result.xml


def test_process_specification_file_project_spec_xlsx_bare_round_diameter(tmp_path):
    """Круглый воздуховод с «голым» диаметром в конце наименования не теряется."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append([
        "Позиция",
        "Наименование и техническая характеристика",
        "Тип, марка, обозначение документа, опросного листа",
        "Код оборудования, изделия, материала",
        "Завод-изготовитель",
        "Единица измерения",
        "Количество",
    ])
    ws.append(["1", "2", "3", "4", "5", "6", "7"])
    ws.append([
        "1",
        'Воздуховод из нержавеющей стали толщиной 0,8 мм класс воздуховодов "П" 200',
        "", "", "Россия", "м", "50",
    ])
    p = tmp_path / "spec.xlsx"
    wb.save(p)

    result = process_specification_file(p.read_bytes(), "spec.xlsx")
    assert "<productRow>" in result.xml
    assert "<material>2</material>" in result.xml  # нержавейка из наименования
    assert "<thickness>0.8</thickness>" in result.xml


def test_process_specification_file_material_matches_source_marker(tmp_path):
    """Материал из наименования исходника не теряется по пути в XML (заказ 844).

    Каждая строка с маркером материала в наименовании должна получить
    соответствующий код: оцинкованная -> 1, нержавеющая -> 2, черная -> 3.
    """
    import re
    import xml.etree.ElementTree as ET
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append([
        "Позиция",
        "Наименование и техническая характеристика",
        "Тип, марка, обозначение документа, опросного листа",
        "Код оборудования, изделия, материала",
        "Завод-изготовитель",
        "Единица измерения",
        "Количество",
    ])
    ws.append(["1", "2", "3", "4", "5", "6", "7"])
    ws.append([
        "1",
        'Воздуховод из оцинкованной стали толщиной согласно СП 60 класс воздуховодов "П" 200',
        "", "", "Россия", "м", "30",
    ])
    ws.append([
        "2",
        'Воздуховод из нержавеющей стали толщиной 0,8 мм класс воздуховодов "П" 200',
        "", "", "Россия", "м", "50",
    ])
    ws.append([
        "3",
        'Отвод круглого воздуховода из черной стали толщиной 0,8 мм D200',
        "", "", "Россия", "шт", "5",
    ])
    p = tmp_path / "spec.xlsx"
    wb.save(p)

    result = process_specification_file(p.read_bytes(), "spec.xlsx")
    rows = ET.fromstring(result.xml).findall(".//productRow")
    assert len(rows) == 3

    expected = {
        "оцинк": "1",
        "нерж": "2",
        "черн": "3",
    }
    for row in rows:
        comment = (row.findtext("comment") or "").lower()
        material = row.findtext("material")
        marker = next(m for m in expected if m in comment)
        assert material == expected[marker], (
            f"материал не соответствует исходнику: {comment[:60]!r} -> {material}"
        )


def test_load_tables_from_pdf_block_fallback(tmp_path):
    """ГОСТ-бланк с многострочными ячейками: find_tables съезжается,
    load_tables_from_pdf должен упасть в разбор текстового слоя блоками."""
    from pathlib import Path

    from api import load_tables_from_pdf

    pdf_path = (
        Path(__file__).parent.parent
        / "examples" / "для обучения 05.09"
        / "ОВ2_СТР_Шипиловский_26-03-02 МОЭК (1).pdf"
    )
    if not pdf_path.exists():
        import pytest

        pytest.skip("PDF проекта Шипиловский не найден в examples")
    result = load_tables_from_pdf(pdf_path.read_bytes())
    assert "block_rows" in result
    rows = result["block_rows"]
    # круглый нерж-воздуховод Ф200, 50 м
    duct = [r for r in rows if r["size"] == "200" and "нержавеющей" in r["name"]]
    assert len(duct) == 1
    assert duct[0]["unit"] == "м"
    assert duct[0]["quantity"] == 50.0
    # секция пожарной защиты 0,9 мм
    fz = [r for r in rows if "0,9 мм" in r["name"]]
    assert fz and any(r["size"] == "150x150" and r["quantity"] == 75.0 for r in fz)
